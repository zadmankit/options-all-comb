"""
استراتژی All Combinations - بازار اختیار معامله ایران
------------------------------------------------------
- جفت‌سازی Call و Put با نماد پایه یکسان و حداکثر اختلاف روز سررسید قابل‌تنظیم (۰ = هم‌سررسید)
- قیمت ورود: po1 (در صورت نبود: pl)
- قیمت پایه: base_pc
- خروجی: All_Comb

نیازمندی‌ها:
    pip install streamlit streamlit-aggrid pandas numpy requests python-dotenv openpyxl
"""

import os
from io import BytesIO
from datetime import datetime

import numpy as np
import pandas as pd
import requests
import streamlit as st
from dotenv import load_dotenv
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from st_aggrid import AgGrid, GridOptionsBuilder, GridUpdateMode, JsCode

# =====================================================================
# تنظیمات اولیه
# =====================================================================
load_dotenv()
DEFAULT_TIMEOUT = 20


def _secret(name: str, default: str = "") -> str:
    """اول st.secrets (Streamlit Cloud)، بعد متغیر محیطی/.env، در نهایت مقدار پیش‌فرض."""
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:  # noqa: BLE001  (وقتی فایل secrets وجود ندارد)
        pass
    return os.getenv(name, default)


API_KEY_DEFAULT = _secret("API_KEY", "")
API_URL_DEFAULT = _secret(
    "API_URL",
    f"https://api.brsapi.ir/Tsetmc/Option.php?key={API_KEY_DEFAULT}" if API_KEY_DEFAULT else "",
)
USER_AGENT_DEFAULT = _secret(
    "USER_AGENT",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36",
)

EXCEL_MAX_ROWS = 1_048_575
ANALYSIS_COLS = [
    "میانگین ریزش", "میانگین مرکزی", "میانگین رشد", "میانگین منفی",
    "میانگین مثبت", "میانگین ساده", "تعداد سود",
]

# ستون‌های اختلاف درصدی نسبت به قیمت پایانی (base_pc)
DIFF_COLS_CALL = ["اختلاف اعمال و پایانی Call (%)", "اختلاف سر به سر و پایانی Call (%)"]
DIFF_COLS_PUT = ["اختلاف اعمال و پایانی Put (%)", "اختلاف سر به سر و پایانی Put (%)"]

st.set_page_config(page_title="استراتژی All Combinations", page_icon="📊", layout="wide")

CUSTOM_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Vazirmatn:wght@300;400;500;700&display=swap');

html, body, .stApp, p, h1, h2, h3, h4, h5, label, li, td, th, input, textarea,
[data-testid="stMetricLabel"] p, [data-testid="stMetricValue"] div {
    font-family: 'Vazirmatn', Tahoma, sans-serif;
}
.stApp, section[data-testid="stSidebar"] { direction: rtl; }
.block-container { padding-top: 1.6rem; padding-bottom: 3rem; max-width: 100%; }

.app-head { border-right: 6px solid #0F766E; padding: 0.2rem 1rem; margin-bottom: 1.2rem; }
.app-head h1 { font-size: 1.7rem; font-weight: 700; margin: 0; padding: 0; color: #14213D; }
.app-head p  { margin: 0.25rem 0 0 0; color: #5B6577; font-size: 0.95rem; }

[data-testid="stMetric"] {
    background: #F4F6F8; border: 1px solid #E1E5EA; border-radius: 8px; padding: 0.7rem 0.9rem;
}
[data-testid="stMetricLabel"] p { font-size: 0.82rem; color: #5B6577; }
[data-testid="stMetricValue"] div { font-size: 1.35rem; font-weight: 700; color: #14213D; }

.stTabs [data-baseweb="tab"] { font-weight: 500; }
.stButton > button, .stDownloadButton > button { border-radius: 8px; font-weight: 500; }
.hint { color: #5B6577; font-size: 0.88rem; }
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)
st.markdown(
    "<div class='app-head'><h1>📊 استراتژی All Combinations</h1>"
    "<p>ترکیب Call و Put هم‌پایه و هم‌سررسید در بازار اختیار معامله ایران</p></div>",
    unsafe_allow_html=True,
)


# =====================================================================
# توابع کمکی و نرمال‌سازی داده
# =====================================================================
def safe_get(d: dict, keys: list, default=None):
    for k in keys:
        if isinstance(d, dict) and k in d and d[k] not in (None, "", " "):
            return d[k]
    return default


def to_num(x):
    if x is None or x == "":
        return None
    try:
        if isinstance(x, str):
            x = x.replace(",", "").strip()
        return float(x)
    except (TypeError, ValueError):
        return None


def detect_option_type(item: dict, symbol) -> str:
    """اول فیلد type، سپس حرف اول نماد (ض = Call ، ط = Put)."""
    t = str(safe_get(item, ["type", "option_type"], "")).strip().lower()
    if t in ("call", "c"):
        return "Call"
    if t in ("put", "p"):
        return "Put"
    if symbol:
        if symbol.startswith("ض"):
            return "Call"
        if symbol.startswith("ط"):
            return "Put"
    return "Unknown"


def pick_price(po1, pl):
    """قیمت ورود: po1 اگر موجود و بزرگ‌تر از صفر بود، وگرنه pl."""
    if po1 is not None and po1 > 0:
        return po1, "po1"
    if pl is not None and pl > 0:
        return pl, "pl"
    return None, None


def normalize_item(item: dict) -> dict:
    symbol = safe_get(item, ["l18", "symbol", "name", "ticker"])
    po1 = to_num(safe_get(item, ["po1"]))
    pl = to_num(safe_get(item, ["pl"]))
    price, src = pick_price(po1, pl)
    return {
        "symbol": symbol or "N/A",
        "option_type": detect_option_type(item, symbol),
        "base_symbol": safe_get(item, ["base_l18", "base_symbol", "underlying_symbol"]),
        "strike": to_num(safe_get(item, ["price_strike", "strike", "exercise_price"])),
        "po1": po1,
        "pl": pl,
        "price": price,
        "price_src": src,
        "base_pc": to_num(safe_get(item, ["base_pc"])),
        "day_remain": to_num(safe_get(item, ["day_remain", "remain_day", "days_remain"])),
        "interest_open": to_num(safe_get(item, ["interest_open", "open_interest"])),
        "tvol": to_num(safe_get(item, ["tvol"])),
    }


def parse_api_response(resp_json) -> pd.DataFrame:
    items = None
    if isinstance(resp_json, list):
        items = resp_json
    elif isinstance(resp_json, dict):
        for k in ("data", "result", "items", "contracts", "rows", "payload"):
            if k in resp_json and isinstance(resp_json[k], list):
                items = resp_json[k]
                break
        if items is None:
            for v in resp_json.values():
                if isinstance(v, list) and v and isinstance(v[0], dict):
                    items = v
                    break
    if not items:
        raise ValueError("پاسخ API حاوی لیست قراردادها نبود.")
    return pd.DataFrame([normalize_item(it) for it in items])


def fetch_api(url: str, user_agent: str, timeout: int = DEFAULT_TIMEOUT) -> pd.DataFrame:
    resp = requests.get(url, headers={"User-Agent": user_agent}, timeout=timeout)
    resp.raise_for_status()
    return parse_api_response(resp.json())


def filter_contracts(df: pd.DataFrame, exclude_zero_price: bool, exclude_zero_oi: bool,
                     exclude_zero_vol: bool, min_days: int):
    """پاک‌سازی و فیلتر قراردادها. خروجی: (DataFrame، آمار)."""
    stats = {"raw": len(df)}
    d = df[df["option_type"].isin(["Call", "Put"])]
    d = d.dropna(subset=["strike", "base_pc", "day_remain", "base_symbol"])
    d = d[(d["strike"] > 0) & (d["base_pc"] > 0)]
    if exclude_zero_price:
        d = d[d["po1"].fillna(0) > 0]  # بدون جایگزین pl
    d = d.dropna(subset=["price"])
    if exclude_zero_oi:
        d = d[d["interest_open"].fillna(0) > 0]
    if exclude_zero_vol:
        d = d[d["tvol"].fillna(0) > 0]
    if min_days > 0:
        d = d[d["day_remain"] >= min_days]
    d = d.reset_index(drop=True)
    stats["calls"] = int((d["option_type"] == "Call").sum())
    stats["puts"] = int((d["option_type"] == "Put").sum())
    stats["pl_share"] = float((d["price_src"] == "pl").mean() * 100) if len(d) else 0.0
    return d, stats


# =====================================================================
# سناریوها و ستون‌های تحلیلی
# =====================================================================
def scenario_cols(range_step: int):
    percents = [int(i * range_step) for i in range(-3, 4)]
    amt_cols = [f"مبلغ {p:+d}%" for p in percents]
    pct_cols = [f"سود/زیان {p:+d}%" for p in percents]
    return percents, amt_cols, pct_cols


def add_analysis(df: pd.DataFrame, pct_cols: list) -> pd.DataFrame:
    """محاسبه برداری ستون‌های تحلیلی (۷ سناریو به‌ترتیب از منفی‌ترین تا مثبت‌ترین)."""
    if df.empty:
        return df
    v = df[pct_cols].to_numpy(float)
    n = len(v)

    pos = np.where(v > 0, v, 0.0)
    neg = np.where(v < 0, v, 0.0)
    n_pos = (v > 0).sum(axis=1)
    n_neg = (v < 0).sum(axis=1)

    out = pd.DataFrame({
        "میانگین ریزش": v[:, 0:3].mean(axis=1),    # سه سناریوی منفی (مثلاً -15، -10، -5)
        "میانگین مرکزی": v[:, 2:5].mean(axis=1),   # سه سناریوی وسط (مثلاً -5، 0، +5)
        "میانگین رشد": v[:, 4:7].mean(axis=1),     # سه سناریوی مثبت (مثلاً +5، +10، +15)
        "میانگین منفی": np.divide(neg.sum(axis=1), n_neg, out=np.zeros(n), where=n_neg > 0),
        "میانگین مثبت": np.divide(pos.sum(axis=1), n_pos, out=np.zeros(n), where=n_pos > 0),
        "میانگین ساده": v.mean(axis=1),
        "تعداد سود": n_pos,
    }, index=df.index).round(2)
    out["تعداد سود"] = out["تعداد سود"].astype(int)
    return pd.concat([df, out], axis=1)


def all_comb_order(range_step: int) -> list:
    """ترتیب ستون‌های All_Comb در رابط کاربری و Excel."""
    _, amt_cols, pct_cols = scenario_cols(range_step)
    return [
        "نماد Call", "قیمت اعمال Call", "قیمت Call",
        "نماد Put", "قیمت اعمال Put", "قیمت Put",
        "نسبت Call:Put",
        "نماد پایه", "روز باقیمانده", "اختلاف روز",
        *pct_cols,
        *ANALYSIS_COLS,
        *amt_cols,
        "قیمت پایه",
        "منبع قیمت Call", "منبع قیمت Put",
        "موقعیت باز Call", *DIFF_COLS_CALL,
        "موقعیت باز Put", *DIFF_COLS_PUT,
        "هزینه کل ورود",
    ]


# =====================================================================
# ساخت All_Comb (هم‌پایه، با حداکثر اختلاف سررسید مجاز)
# =====================================================================
def build_all_comb(contracts: pd.DataFrame, range_step: int, ratio: float, max_day_gap: int = 0):
    percents, amt_cols, pct_cols = scenario_cols(range_step)
    pcts = np.array(percents, float)
    ratios = list(dict.fromkeys([(1.0, 1.0), (float(ratio), 1.0), (1.0, float(ratio))]))

    calls = contracts[contracts["option_type"] == "Call"]
    puts = contracts[contracts["option_type"] == "Put"]
    put_groups = {k: g for k, g in puts.groupby("base_symbol")}

    frames, n_groups = [], 0
    for base_sym, cg in calls.groupby("base_symbol"):
        pg = put_groups.get(base_sym)
        if pg is None:
            continue
        ci, pi = np.meshgrid(np.arange(len(cg)), np.arange(len(pg)), indexing="ij")
        ci, pi = ci.ravel(), pi.ravel()
        dc = cg["day_remain"].to_numpy(float)[ci]
        dp = pg["day_remain"].to_numpy(float)[pi]
        ok = np.abs(dc - dp) <= max_day_gap          # «کمتر یا مساوی»
        if not ok.any():
            continue
        ci, pi, dc, dp = ci[ok], pi[ok], dc[ok], dp[ok]
        n_groups += 1
        c = cg.iloc[ci].reset_index(drop=True)
        p = pg.iloc[pi].reset_index(drop=True)
        days_near = np.minimum(dc, dp).astype(int)   # سررسید نزدیک‌تر
        days_gap = np.abs(dc - dp).astype(int)

        base = c["base_pc"].to_numpy(float)
        kc, kp = c["strike"].to_numpy(float), p["strike"].to_numpy(float)
        pc, pp = c["price"].to_numpy(float), p["price"].to_numpy(float)
        future = base[:, None] * (1 + pcts[None, :] / 100)

        # قیمت سر به سر: Call = اعمال + قیمت ، Put = اعمال − قیمت
        be_c, be_p = kc + pc, kp - pp
        # اختلاف درصدی = (قیمت پایانی − سطح) / قیمت پایانی × ۱۰۰
        d_k_c = (base - kc) / base * 100
        d_be_c = (base - be_c) / base * 100
        d_k_p = (base - kp) / base * 100
        d_be_p = (base - be_p) / base * 100

        for rc, rp in ratios:
            cost = rc * pc + rp * pp
            value = np.maximum(0, future - kc[:, None]) * rc + np.maximum(0, kp[:, None] - future) * rp
            pnl = value - cost[:, None]
            pct = pnl / cost[:, None] * 100

            data = {
                "نماد پایه": base_sym,
                "روز باقیمانده": days_near,
                "اختلاف روز": days_gap,
                "قیمت پایه": base,
                "نماد Call": c["symbol"].to_numpy(),
                "قیمت اعمال Call": kc,
                "قیمت Call": pc,
                "منبع قیمت Call": c["price_src"].to_numpy(),
                "موقعیت باز Call": c["interest_open"].fillna(0).astype(int).to_numpy(),
                DIFF_COLS_CALL[0]: np.round(d_k_c, 2),
                DIFF_COLS_CALL[1]: np.round(d_be_c, 2),
                "نماد Put": p["symbol"].to_numpy(),
                "قیمت اعمال Put": kp,
                "قیمت Put": pp,
                "منبع قیمت Put": p["price_src"].to_numpy(),
                "موقعیت باز Put": p["interest_open"].fillna(0).astype(int).to_numpy(),
                DIFF_COLS_PUT[0]: np.round(d_k_p, 2),
                DIFF_COLS_PUT[1]: np.round(d_be_p, 2),
                "نسبت Call:Put": f"{rc:g}:{rp:g}",
                "هزینه کل ورود": np.round(cost, 2),
            }
            for k, col in enumerate(amt_cols):
                data[col] = np.round(pnl[:, k], 2)
            for k, col in enumerate(pct_cols):
                data[col] = np.round(pct[:, k], 2)
            frames.append(pd.DataFrame(data))

    if not frames:
        return pd.DataFrame(), 0

    df = pd.concat(frames, ignore_index=True)
    df = add_analysis(df, pct_cols)
    df = df[all_comb_order(range_step)]
    return df, n_groups


# =====================================================================
# پایپ‌لاین اصلی (ذخیره نتایج در session_state)
# =====================================================================
def run_pipeline(raw_df: pd.DataFrame, p: dict):
    with st.status("در حال پردازش…", expanded=True) as status:
        st.write("فیلتر قراردادها")
        contracts, stats = filter_contracts(
            raw_df, p["exclude_zero_price"], p["exclude_zero_oi"], p["exclude_zero_vol"], p["min_days"]
        )
        if contracts.empty:
            status.update(label="قراردادی باقی نماند", state="error")
            st.session_state["results"] = None
            st.error("پس از اعمال فیلترها هیچ قراردادی باقی نماند. فیلترها را بازبینی کنید.")
            return

        st.write("ساخت All_Comb")
        df_all, n_groups = build_all_comb(contracts, p["range_step"], p["ratio"], p["max_day_gap"])
        status.update(label="پردازش کامل شد", state="complete", expanded=False)

    st.session_state["results"] = {
        "contracts": contracts, "stats": stats, "groups": n_groups,
        "all": df_all, "range_step": p["range_step"], "ts": datetime.now(),
    }
    st.session_state["excel_bytes"] = None


# =====================================================================
# نمایش: جدول، فیلتر، نمودار
# =====================================================================
HEAT_JS = JsCode("""
function(params) {
    var v = params.value;
    if (v === null || v === undefined || isNaN(v)) { return null; }
    var a = Math.min(Math.abs(v) / 150, 1) * 0.55 + 0.08;
    if (v > 0) { return {'backgroundColor': 'rgba(21,128,61,' + a + ')', 'fontWeight': '600'}; }
    if (v < 0) { return {'backgroundColor': 'rgba(185,28,28,' + a + ')'}; }
    return null;
}
""")


def show_grid(df: pd.DataFrame, key: str, height: int = 540):
    gb = GridOptionsBuilder.from_dataframe(df)
    gb.configure_default_column(filter=True, sortable=True, resizable=True, minWidth=120)
    gb.configure_pagination(enabled=True, paginationPageSize=25)
    gb.configure_grid_options(enableRtl=True)
    heat_cols = ["میانگین ریزش", "میانگین مرکزی", "میانگین رشد", "میانگین ساده", "میانگین مثبت", "میانگین منفی"]
    for col in df.columns:
        if col.startswith("سود/زیان") or col in heat_cols:
            gb.configure_column(col, cellStyle=HEAT_JS)
    AgGrid(
        df,
        gridOptions=gb.build(),
        allow_unsafe_jscode=True,
        enable_enterprise_modules=False,
        update_mode=GridUpdateMode.NO_UPDATE,
        fit_columns_on_grid_load=False,
        theme="alpine",
        height=height,
        key=key,
    )


def filtered_view(df: pd.DataFrame, key: str):
    """فیلترهای سریع بالای جدول. خروجی: (df نمایشی، تعداد کل مطابق فیلتر)."""
    cols = st.columns(3)
    bases = cols[0].multiselect("نماد پایه", sorted(df["نماد پایه"].unique().tolist()), key=f"{key}_base")
    days = cols[1].multiselect("روز باقیمانده", sorted(df["روز باقیمانده"].unique().tolist()), key=f"{key}_days")
    sort_by = cols[2].selectbox(
        "مرتب‌سازی بر اساس", ANALYSIS_COLS, index=ANALYSIS_COLS.index("میانگین ساده"), key=f"{key}_sort"
    )

    c1, c2, _ = st.columns([1, 1, 3])
    max_rows = c1.number_input("حداکثر ردیف نمایشی", 100, 100_000, 5_000, 500, key=f"{key}_max")
    show_amounts = c2.checkbox("نمایش ستون‌های مبلغ", value=False, key=f"{key}_amt")

    m = pd.Series(True, index=df.index)
    if bases:
        m &= df["نماد پایه"].isin(bases)
    if days:
        m &= df["روز باقیمانده"].isin(days)
    d = df[m].sort_values(sort_by, ascending=False)
    total = len(d)
    d = d.head(int(max_rows))
    if not show_amounts:
        d = d.drop(columns=[c for c in d.columns if c.startswith("مبلغ")])
    return d, total


# =====================================================================
# Excel
# =====================================================================
def build_excel(res: dict) -> bytes:
    buf = BytesIO()
    header_fill = PatternFill("solid", fgColor="14213D")
    header_font = Font(bold=True, color="FFFFFF")
    green, red = Font(color="15803D", bold=True), Font(color="B91C1C")

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in (("All_Comb", res["all"]),):
            if df is None or df.empty:
                continue
            d = df.iloc[:EXCEL_MAX_ROWS]
            d.to_excel(writer, index=False, sheet_name=name)
            ws = writer.sheets[name]
            ws.sheet_view.rightToLeft = True
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(d.columns, start=1):
                cell = ws.cell(row=1, column=i)
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.column_dimensions[get_column_letter(i)].width = max(12, min(38, len(str(col)) + 6))
                if str(col).startswith("سود/زیان"):
                    rng = f"{get_column_letter(i)}2:{get_column_letter(i)}{len(d) + 1}"
                    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=green))
                    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=red))
    return buf.getvalue()


# =====================================================================
# سایدبار
# =====================================================================
sb = st.sidebar
sb.markdown("## ⚙️ تنظیمات")

with sb.expander("🔌 اتصال API"):
    api_url = st.text_input("API URL", API_URL_DEFAULT)
    api_key = st.text_input("API Key", API_KEY_DEFAULT, type="password")
    user_agent = st.text_input("User-Agent", USER_AGENT_DEFAULT)

with sb.expander("📐 سناریوها", expanded=True):
    range_step = st.number_input(
        "بازه درصدی سناریوها", 1, 50, 5, 1,
        help="مثلاً ۵ ← سناریوها: ‎−15، −10، −5، 0، +5، +10، +15 درصد",
    )
    ratio = st.number_input(
        "ضریب نسبت", 0.5, 10.0, 2.0, 0.5,
        help="نسبت‌های ساخته‌شده: 1:1 ، ضریب:1 ، 1:ضریب",
    )

with sb.expander("🧹 فیلتر قراردادها"):
    exclude_zero_price = st.checkbox(
        "حذف قیمت صفر", False,
        help="قراردادهایی که po1 صفر دارند حذف می‌شوند (جایگزین pl غیرفعال می‌شود).",
    )
    exclude_zero_oi = st.checkbox("حذف موقعیت باز صفر", False)
    exclude_zero_vol = st.checkbox(
        "حذف حجم معاملات صفر", True,
        help="قراردادهایی که امروز هیچ معامله‌ای نداشته‌اند (tvol صفر) حذف می‌شوند.",
    )
    min_days = st.number_input("حداقل روز باقیمانده", 0, 365, 0, 1)

with sb.expander("📅 اختلاف سررسید", expanded=True):
    max_day_gap = st.number_input(
        "حداکثر اختلاف روز سررسید", 0, 365, 0, 1,
        help="Call و Put با نماد پایه یکسان جفت می‌شوند اگر اختلاف روز باقیمانده‌شان کمتر یا مساوی این عدد باشد. "
             "۰ یعنی فقط هم‌سررسید.",
    )

sb.markdown("---")
fetch_btn = sb.button("📥 دریافت اطلاعات و محاسبه", type="primary", use_container_width=True)
recalc_btn = sb.button(
    "🔄 محاسبه مجدد با داده فعلی", use_container_width=True,
    help="بدون دریافت دوباره از API، با تنظیمات جدید دوباره محاسبه می‌کند.",
)

params = {
    "range_step": int(range_step), "ratio": float(ratio),
    "exclude_zero_price": exclude_zero_price, "exclude_zero_oi": exclude_zero_oi,
    "exclude_zero_vol": exclude_zero_vol, "min_days": int(min_days),
    "max_day_gap": int(max_day_gap),
}

# =====================================================================
# اجرا
# =====================================================================
try:
    if fetch_btn:
        if not api_url:
            st.error("API URL را وارد کنید.")
        else:
            url = api_url
            if api_key and "key=" not in url:
                url += ("&" if "?" in url else "?") + f"key={api_key}"
            with st.spinner("در حال دریافت داده از API…"):
                st.session_state["raw_df"] = fetch_api(url, user_agent)
                st.session_state["fetched_at"] = datetime.now()
            run_pipeline(st.session_state["raw_df"], params)
    elif recalc_btn:
        if st.session_state.get("raw_df") is None:
            st.warning("ابتدا یک‌بار «دریافت اطلاعات و محاسبه» را بزنید.")
        else:
            run_pipeline(st.session_state["raw_df"], params)
except requests.HTTPError as e:
    code = e.response.status_code if e.response is not None else "?"
    st.error(f"خطای API (کد {code}). کلید یا آدرس را بررسی کنید.")
except Exception as e:  # noqa: BLE001
    st.error(f"خطا در دریافت یا پردازش داده‌ها: {e}")

# =====================================================================
# نمایش نتایج
# =====================================================================
res = st.session_state.get("results")

if res is None:
    st.info("برای شروع، از سایدبار روی «دریافت اطلاعات و محاسبه» بزنید.")
    st.stop()

df_all, stats = res["all"], res["stats"]

fetched_at = st.session_state.get("fetched_at")
_fetched_txt = f"آخرین دریافت داده: {fetched_at:%Y-%m-%d %H:%M:%S} | " if fetched_at else ""
st.caption(f"{_fetched_txt}آخرین محاسبه: {res['ts']:%H:%M:%S}")

m = st.columns(5)
m[0].metric("قرارداد Call", f"{stats['calls']:,}")
m[1].metric("قرارداد Put", f"{stats['puts']:,}")
m[2].metric("نماد پایه دارای جفت", f"{res['groups']:,}")
m[3].metric("ردیف All_Comb", f"{len(df_all):,}")
m[4].metric("بهترین میانگین ساده", f"{df_all['میانگین ساده'].max():.1f}%" if not df_all.empty else "—")

if stats["pl_share"] > 0:
    st.warning(
        f"{stats['pl_share']:.0f}٪ قراردادها به‌جای po1 با قیمت آخرین معامله (pl) قیمت‌گذاری شده‌اند. "
        "ممکن است قابل‌اجرا نباشند؛ ستون «منبع قیمت» را ببینید."
    )
if df_all.empty:
    st.warning("هیچ جفت Call/Put هم‌پایه‌ای با اختلاف سررسید مجاز پیدا نشد.")

tab_all, tab_dl = st.tabs(["🔗 All_Comb", "📤 دانلود"])

with tab_all:
    if df_all.empty:
        st.info("داده‌ای برای نمایش نیست.")
    else:
        view, total = filtered_view(df_all, "all")
        st.caption(f"نمایش {len(view):,} ردیف از {total:,} ردیف مطابق فیلتر")
        show_grid(view, f"grid_all_{res['ts']:%H%M%S%f}")

with tab_dl:
    st.markdown("##### خروجی Excel (شیت All_Comb)")
    if df_all.empty:
        st.info("داده‌ای برای خروجی نیست.")
    else:
        if len(df_all) > EXCEL_MAX_ROWS:
            st.warning("تعداد ردیف‌ها از سقف Excel بیشتر است؛ فقط ۱٬۰۴۸٬۵۷۵ ردیف اول نوشته می‌شود.")
        if len(df_all) > 300_000:
            st.caption("ساخت فایل برای حجم بالا ممکن است چند دقیقه طول بکشد.")
        if st.button("🛠️ ساخت فایل Excel"):
            with st.spinner("در حال ساخت فایل Excel…"):
                st.session_state["excel_bytes"] = build_excel(res)
        data = st.session_state.get("excel_bytes")
        if data:
            st.download_button(
                f"📥 دانلود Excel ({len(data) / 1024 / 1024:.1f} MB)",
                data=data,
                file_name=f"all_combinations_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )

st.markdown("---")
st.markdown(
    "<div style='text-align:center;color:#8A93A3;font-size:0.85rem'>All Combinations — نسخه ۱۴</div>",
    unsafe_allow_html=True,
)